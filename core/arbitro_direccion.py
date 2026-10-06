"""Árbitro de dirección (B): corrige la dirección de escritura en los casos ambiguos.

El escáner elige "derecha" apenas la celda vecina de la derecha está vacía. En una tabla con encabezados
(p. ej. "Número de Cuenta" sobre su celda de captura) la celda vacía de la derecha puede estar **fuera de
la tabla**, mientras que la de captura real está debajo, con el mismo ancho combinado y con bordes.

Este paso solo evalúa los casos ambiguos (derecha y abajo vacías) y solo puede cambiar "derecha" por
"abajo". Exige siempre dos condiciones, y además una ventaja clara de puntaje:

  * condición 1: la celda de abajo tiene el mismo ancho que el rótulo,
  * condición 2: la celda de abajo es una caja de captura (con bordes en al menos tres lados),
  * ventaja: pesan a favor la celda de la derecha fuera de la tabla (sin bordes ni continuación), el relleno
    de captura abajo y la consistencia con los demás rótulos de la fila (encabezados cuyo valor va debajo).

Un relleno sin caja debajo (títulos de sección, rótulos con opciones a la derecha) nunca basta.

El resultado se anota en el elemento (``direccionArbitrada``); las etapas posteriores lo respetan.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import range_boundaries

UMBRAL_DE_DECISION = 3
LADOS = ("top", "bottom", "left", "right")


@dataclass(frozen=True)
class CambioDeDireccion:
    """Un rótulo cuya dirección pasó de "derecha" a "abajo", con la evidencia que lo justificó."""

    hoja: str
    fila: int
    columna: int
    rotulo: str
    puntaje_abajo: int
    puntaje_derecha: int
    motivos: Tuple[str, ...]


def _rango_de(hoja: Any, fila: int, columna: int) -> Tuple[int, int, int, int]:
    """(min_fila, min_col, max_fila, max_col) del rango combinado que contiene la celda, o la celda sola."""
    for rango in hoja.merged_cells.ranges:
        if rango.min_row <= fila <= rango.max_row and rango.min_col <= columna <= rango.max_col:
            return rango.min_row, rango.min_col, rango.max_row, rango.max_col
    return fila, columna, fila, columna


def _bordes(hoja: Any, rango: Tuple[int, int, int, int]) -> int:
    """Cuántos lados del perímetro del rango tienen borde (0-4)."""
    f0, c0, f1, c1 = rango
    lados = {
        "top": [hoja.cell(row=f0, column=c) for c in range(c0, c1 + 1)],
        "bottom": [hoja.cell(row=f1, column=c) for c in range(c0, c1 + 1)],
        "left": [hoja.cell(row=f, column=c0) for f in range(f0, f1 + 1)],
        "right": [hoja.cell(row=f, column=c1) for f in range(f0, f1 + 1)],
    }
    total = 0
    for lado, celdas in lados.items():
        if any(getattr(celda.border, lado) is not None and getattr(celda.border, lado).style for celda in celdas):
            total += 1
    return total


def _lado_con_borde(hoja: Any, rango: Tuple[int, int, int, int], lado: str) -> bool:
    """True si algún tramo del lado indicado (top/bottom/left/right) del rango tiene borde."""
    f0, c0, f1, c1 = rango
    celdas = {
        "top": [hoja.cell(row=f0, column=c) for c in range(c0, c1 + 1)],
        "bottom": [hoja.cell(row=f1, column=c) for c in range(c0, c1 + 1)],
        "left": [hoja.cell(row=f, column=c0) for f in range(f0, f1 + 1)],
        "right": [hoja.cell(row=f, column=c1) for f in range(f0, f1 + 1)],
    }[lado]
    return any(getattr(c.border, lado) is not None and getattr(c.border, lado).style for c in celdas)


def _con_relleno(hoja: Any, rango: Tuple[int, int, int, int]) -> bool:
    celda = hoja.cell(row=rango[0], column=rango[1])
    return bool(celda.fill is not None and celda.fill.fill_type)


def _hay_bordes_mas_alla(hoja: Any, fila: int, columna: int) -> bool:
    """True si la celda o las dos siguientes de la derecha tienen algún borde (la tabla continúa)."""
    for c in range(columna, columna + 3):
        celda = hoja.cell(row=fila, column=c)
        if any(getattr(celda.border, lado) is not None and getattr(celda.border, lado).style for lado in LADOS):
            return True
    return False


def _direccion_de_hermano(elemento: Dict[str, Any]) -> Optional[str]:
    """Dirección que el escáner dejó clara para un rótulo vecino (solo los casos no ambiguos)."""
    derecha_vacia, abajo_vacia = bool(elemento.get("derechaVacia")), bool(elemento.get("abajoVacia"))
    if abajo_vacia and not derecha_vacia:
        return "abajo"
    if derecha_vacia and not abajo_vacia:
        return "derecha"
    return None


def _evaluar(hoja: Any, elemento: Dict[str, Any], hermanos: List[Dict[str, Any]]) -> Tuple[int, int, List[str], bool]:
    fila, columna = int(elemento["fila"]), int(elemento["columna"])
    if elemento.get("coordMerge"):
        c0, f0, c1, f1 = range_boundaries(str(elemento["coordMerge"]))
    else:
        c0, f0, c1, f1 = columna, fila, columna, fila
    derecha = _rango_de(hoja, f0, c1 + 1)
    abajo = _rango_de(hoja, f1 + 1, c0)

    motivos: List[str] = []
    abajo_puntos, derecha_puntos = 0, 0

    # Celda de abajo: mismo ancho que el rótulo, con bordes y con relleno.
    mismo_ancho = abajo[1] == c0 and abajo[3] == c1
    if mismo_ancho:
        abajo_puntos += 2
        motivos.append("la celda de abajo tiene el mismo ancho que el rótulo")
    bordes_abajo = _bordes(hoja, abajo)
    # El borde inferior del encabezado es el borde superior de la celda de captura: se comparte.
    if bordes_abajo < 4 and _lado_con_borde(hoja, (f0, c0, f1, c1), "bottom") and not _lado_con_borde(hoja, abajo, "top"):
        bordes_abajo += 1
    es_caja = bordes_abajo >= 3
    if bordes_abajo >= 3:
        abajo_puntos += 2
        motivos.append("la celda de abajo es una caja con bordes")
    elif bordes_abajo >= 1:
        abajo_puntos += 1
    if _con_relleno(hoja, abajo):
        abajo_puntos += 1
        motivos.append("la celda de abajo tiene relleno de captura")

    # Celda de la derecha: caja propia, o fuera de la tabla (sin bordes y sin continuación).
    bordes_derecha = _bordes(hoja, derecha)
    if bordes_derecha >= 3:
        derecha_puntos += 2
    elif bordes_derecha >= 1:
        derecha_puntos += 1
    elif not _hay_bordes_mas_alla(hoja, derecha[0], derecha[1]) and bordes_abajo >= 1:
        derecha_puntos -= 2
        motivos.append("la celda de la derecha está fuera de la tabla (sin bordes)")
    if _con_relleno(hoja, derecha):
        derecha_puntos += 1

    # Consistencia con los demás rótulos de la misma fila.
    direcciones = [d for d in (_direccion_de_hermano(h) for h in hermanos) if d]
    if direcciones and all(d == "abajo" for d in direcciones):
        abajo_puntos += 2
        motivos.append(f"los {len(direcciones)} rótulos vecinos de la fila escriben hacia abajo")
    elif direcciones and all(d == "derecha" for d in direcciones):
        derecha_puntos += 2
    return abajo_puntos, derecha_puntos, motivos, mismo_ancho and es_caja


def arbitrar_direcciones(elementos: List[Dict[str, Any]], archivo_bytes: bytes) -> List[CambioDeDireccion]:
    """Marca como "abajo" los rótulos ambiguos cuya evidencia visual lo respalda; devuelve los cambios."""
    candidatos = [
        e for e in elementos
        if e.get("derechaVacia") and e.get("abajoVacia") and not e.get("esCasillaVerificacion")
        and not e.get("direccionArbitrada")
    ]
    if not candidatos:
        return []
    libro = openpyxl.load_workbook(BytesIO(archivo_bytes), data_only=False)
    por_fila: Dict[Tuple[str, int], List[Dict[str, Any]]] = {}
    for e in elementos:
        por_fila.setdefault((str(e.get("hoja", "")), int(e.get("fila", 0) or 0)), []).append(e)

    cambios: List[CambioDeDireccion] = []
    for elemento in candidatos:
        nombre_hoja = str(elemento.get("hoja", ""))
        if nombre_hoja not in libro.sheetnames:
            continue
        hermanos = [h for h in por_fila[(nombre_hoja, int(elemento["fila"]))] if h is not elemento]
        try:
            abajo, derecha, motivos, requisitos = _evaluar(libro[nombre_hoja], elemento, hermanos)
        except (KeyError, ValueError, TypeError):
            continue
        if requisitos and abajo - derecha >= UMBRAL_DE_DECISION:
            elemento["direccionArbitrada"] = "abajo"
            elemento["direccionArbitradaMotivo"] = "; ".join(motivos)
            cambios.append(CambioDeDireccion(
                nombre_hoja, int(elemento["fila"]), int(elemento["columna"]), str(elemento.get("valor") or ""),
                abajo, derecha, tuple(motivos),
            ))
    return cambios
