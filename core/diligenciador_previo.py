"""Detección de datos de un diligenciador anterior ya escritos en un formulario.

Un formulario generado antes por la app (o preparado a mano) puede traer los datos de otra persona en
el bloque de "Diligenciado por / Contacto comercial". Como ya están llenos, el escritor los respeta y
el cambio de perfil no tendría efecto. Este módulo los reconoce **por su valor**: una celda cuyo
contenido es el nombre, correo, cédula o teléfono de un diligenciador conocido (distinto del activo)
es un dato variable del diligenciador y se reemplaza por el del perfil activo.

Decisiones de seguridad:
  * Solo se tocan valores que identifican de forma específica a otro diligenciador conocido; no se
    infiere nada por la posición ni por el rótulo.
  * Se excluyen los valores que también son datos fijos de la empresa o del representante legal
    (p. ej. un correo corporativo compartido), para no alterar datos que no son del diligenciador.
  * Un nombre completo basta para reconocer al diligenciador anterior; un correo, una cédula o un
    teléfono solos no (podrían ser un dato fijo de la plantilla, como un correo de facturación), así que
    exigen otra celda de identidad del mismo diligenciador a pocas filas.
  * Cargo, dirección y ciudad solo se reemplazan cuando están junto a una celda de identidad del
    mismo diligenciador anterior (son valores demasiado genéricos para reconocerlos por sí solos).
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from io import BytesIO
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

import openpyxl

from core.profile_manager import aplanar_perfil

IDENTIDAD = {
    "nombre": "responsable_nombre",
    "correo": "responsable_correo",
    "cedula": "responsable_cedula",
    "telefono": "responsable_telefono",
}
COMPLEMENTARIOS = {
    "cargo": "responsable_cargo",
    "direccion": "responsable_direccion",
    "ciudad": "responsable_ciudad",
}
FILAS_DE_CERCANIA = 3
MIN_DIGITOS_CEDULA = 5
MIN_DIGITOS_TELEFONO = 7
MIN_LONGITUD_NOMBRE = 6
MIN_LONGITUD_COMPLEMENTARIO = 3

MOTIVO = "Dato de un diligenciador anterior ({previo}): se reemplaza por el del perfil activo"


@dataclass(frozen=True)
class ReemplazoDiligenciador:
    """Una celda que contiene un dato de otro diligenciador y debe pasar al perfil activo."""

    hoja: str
    fila: int
    columna: int
    campo: str
    valor_anterior: str
    valor_nuevo: str
    operador_previo: str
    rotulo_cercano: str = ""


def _normalizar(valor: Any) -> str:
    sin_tildes = "".join(
        c for c in unicodedata.normalize("NFD", str(valor or "").lower()) if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"[^a-z0-9@._-]+", " ", sin_tildes).strip()


def _digitos(valor: Any) -> str:
    return re.sub(r"\D", "", str(valor or ""))


def _huellas(operador: Dict[str, Any]) -> Dict[str, str]:
    """Valores que identifican a un diligenciador, ya normalizados (campo -> huella)."""
    resultado: Dict[str, str] = {}
    nombre = _normalizar(operador.get("nombre"))
    if len(nombre) >= MIN_LONGITUD_NOMBRE and len(nombre.split()) >= 2:
        resultado["nombre"] = nombre
    correo = _normalizar(operador.get("correo"))
    if "@" in correo:
        resultado["correo"] = correo
    cedula = _digitos(operador.get("cedula"))
    if len(cedula) >= MIN_DIGITOS_CEDULA:
        resultado["cedula"] = cedula
    telefono = _digitos(operador.get("telefono") or operador.get("celular"))
    if len(telefono) >= MIN_DIGITOS_TELEFONO:
        resultado["telefono"] = telefono
    return resultado


def _valores_protegidos(datos_empresa: Dict[str, Any]) -> Tuple[Set[str], Set[str]]:
    """Textos y dígitos de los datos fijos de la empresa (todo salvo los campos del diligenciador)."""
    textos: Set[str] = set()
    digitos: Set[str] = set()
    for clave, valor in aplanar_perfil(datos_empresa).items():
        if clave.startswith("responsable_") or clave == "operador" or isinstance(valor, (dict, list, bool)) or valor is None:
            continue
        textos.add(_normalizar(valor))
        d = _digitos(valor)
        if len(d) >= MIN_DIGITOS_CEDULA:
            digitos.add(d)
    textos.discard("")
    return textos, digitos


def _huellas_previas(
    operadores_conocidos: Iterable[Dict[str, Any]],
    activo: Dict[str, Any],
    protegidos: Tuple[Set[str], Set[str]],
) -> List[Tuple[Dict[str, Any], Dict[str, str]]]:
    """Diligenciadores anteriores con las huellas que no coinciden con las del activo ni con la empresa."""
    activo_huellas = set(_huellas(activo).values())
    textos_protegidos, digitos_protegidos = protegidos
    previos: List[Tuple[Dict[str, Any], Dict[str, str]]] = []
    vistos: Set[Tuple[Tuple[str, str], ...]] = set()
    for operador in operadores_conocidos:
        huellas = {
            campo: huella for campo, huella in _huellas(operador).items()
            if huella not in activo_huellas and huella not in textos_protegidos and huella not in digitos_protegidos
        }
        clave = tuple(sorted(huellas.items()))
        if huellas and clave not in vistos:
            vistos.add(clave)
            previos.append((operador, huellas))
    return previos


def _huella_de_celda(campo: str, valor: Any) -> str:
    if campo in ("cedula", "telefono"):
        return _digitos(valor)
    return _normalizar(valor)


def _rotulo_a_la_izquierda(hoja: Any, fila: int, columna: int) -> str:
    for c in range(columna - 1, max(columna - 8, 0), -1):
        valor = hoja.cell(row=fila, column=c).value
        if isinstance(valor, str) and valor.strip() and not valor.strip().startswith("="):
            return valor.strip()[:60]
    return ""


def detectar_reemplazos(
    archivo_bytes: bytes,
    nuevo: Dict[str, Any],
    operadores_conocidos: Iterable[Dict[str, Any]],
    datos_empresa: Optional[Dict[str, Any]] = None,
) -> List[ReemplazoDiligenciador]:
    """Celdas del libro con datos de otro diligenciador conocido, y el valor que deben tomar.

    ``nuevo`` usa las claves ``responsable_*`` del perfil activo (como en los datos fusionados).
    Si el perfil activo no tiene un dato, el reemplazo trae ``valor_nuevo`` vacío: el dato anterior
    se retira en vez de dejar el de otra persona.
    """
    activo = {
        "nombre": nuevo.get("responsable_nombre"), "correo": nuevo.get("responsable_correo"),
        "cedula": nuevo.get("responsable_cedula"),
        "telefono": nuevo.get("responsable_telefono") or nuevo.get("responsable_celular"),
    }
    protegidos = _valores_protegidos(datos_empresa or {})
    previos = _huellas_previas(operadores_conocidos, activo, protegidos)
    if not previos:
        return []

    libro = openpyxl.load_workbook(BytesIO(archivo_bytes), data_only=False)
    reemplazos: List[ReemplazoDiligenciador] = []

    for hoja in libro.worksheets:
        celdas: List[Tuple[int, int, Any]] = [
            (celda.row, celda.column, celda.value)
            for fila in hoja.iter_rows()
            for celda in fila
            if celda.value is not None and not (isinstance(celda.value, str) and celda.value.strip().startswith("="))
        ]
        for operador, huellas in previos:
            encontradas: List[Tuple[int, int, str, Any]] = []
            for fila, columna, valor in celdas:
                for campo, huella in huellas.items():
                    if _huella_de_celda(campo, valor) == huella:
                        encontradas.append((fila, columna, campo, valor))
            encontradas = [
                hit for hit in encontradas
                if hit[2] == "nombre" or any(
                    otro is not hit and abs(otro[0] - hit[0]) <= FILAS_DE_CERCANIA for otro in encontradas
                )
            ]
            if not encontradas:
                continue

            ocupadas = {(f, c) for f, c, _, _ in encontradas}
            filas_ancla = [f for f, _, _, _ in encontradas]
            for fila, columna, campo, valor in encontradas:
                reemplazos.append(_reemplazo(hoja, fila, columna, IDENTIDAD[campo], valor, nuevo, operador))

            for campo_comp, clave_perfil in COMPLEMENTARIOS.items():
                anterior = _normalizar(operador.get(campo_comp))
                if len(anterior) < MIN_LONGITUD_COMPLEMENTARIO:
                    continue
                if anterior == _normalizar(nuevo.get(clave_perfil)):
                    continue
                if anterior in protegidos[0]:
                    continue
                for fila, columna, valor in celdas:
                    if (fila, columna) in ocupadas or _normalizar(valor) != anterior:
                        continue
                    if any(abs(fila - ancla) <= FILAS_DE_CERCANIA for ancla in filas_ancla):
                        reemplazos.append(_reemplazo(hoja, fila, columna, clave_perfil, valor, nuevo, operador))
    return reemplazos


def _reemplazo(
    hoja: Any, fila: int, columna: int, campo: str, valor: Any, nuevo: Dict[str, Any], operador: Dict[str, Any]
) -> ReemplazoDiligenciador:
    valor_nuevo = nuevo.get(campo)
    if campo == "responsable_telefono" and not valor_nuevo:
        valor_nuevo = nuevo.get("responsable_celular")
    return ReemplazoDiligenciador(
        hoja=hoja.title, fila=fila, columna=columna, campo=campo, valor_anterior=str(valor).strip(),
        valor_nuevo=str(valor_nuevo or "").strip(), operador_previo=str(operador.get("nombre") or "otro diligenciador"),
        rotulo_cercano=_rotulo_a_la_izquierda(hoja, fila, columna),
    )


def _destino(item: Dict[str, Any]) -> Optional[Tuple[str, int, int]]:
    """Celda destino de un ítem del plan (misma regla que usa el escritor)."""
    hoja = str(item.get("hoja") or "")
    fila, columna = item.get("fila_destino"), item.get("columna_destino")
    if (not fila or not columna) and item.get("fila") and item.get("columna"):
        f, c = int(item["fila"]), int(item["columna"])
        ubicacion = str(item.get("ubicacion") or "derecha").lower()
        fila, columna = (f + 1, c) if ubicacion == "abajo" else ((f, c) if ubicacion == "misma" else (f, c + 1))
    return (hoja, int(fila), int(columna)) if hoja and fila and columna else None


def aplicar_al_plan(plan: List[Dict[str, Any]], reemplazos: Iterable[ReemplazoDiligenciador]) -> Tuple[List[Dict[str, Any]], int]:
    """Agrega al plan (o marca en él) los reemplazos; devuelve el plan y cuántos se aplicaron.

    Si la celda ya tiene una directiva del diligenciador (``responsable_*``) solo se marca para que el
    escritor pueda sobrescribir el dato previo; si tiene una directiva de otro dominio no se toca; si no
    tiene ninguna, se agrega una nueva.
    """
    resultado = list(plan)
    por_destino = {d: i for i, item in enumerate(resultado) if (d := _destino(item)) is not None}
    aplicados = 0
    for r in reemplazos:
        indice = por_destino.get((r.hoja, r.fila, r.columna))
        if indice is not None:
            existente = resultado[indice]
            if str(existente.get("campo") or "").startswith("responsable_"):
                resultado[indice] = {
                    **existente, "sobrescribir_valor_previo": True, "valor_anterior": r.valor_anterior,
                    "motivo": MOTIVO.format(previo=r.operador_previo),
                }
                aplicados += 1
            continue
        resultado.append({
            "hoja": r.hoja, "fila": r.fila, "columna": r.columna, "fila_destino": r.fila, "columna_destino": r.columna,
            "rotulo": r.rotulo_cercano, "rotulo_original": r.rotulo_cercano, "ubicacion": "misma", "campo": r.campo,
            "valor": None, "valor_a_escribir": r.valor_nuevo, "valor_anterior": r.valor_anterior,
            "seccion": "Diligenciado por (datos previos)", "tipo_elemento": "FIELD", "estado": "APROBADO",
            "nivel_confianza": "ALTA", "confianza_score": 0.9, "motivo": MOTIVO.format(previo=r.operador_previo),
            "requiereMerge": False, "celdasAMergear": 1, "anchoLinea": 1, "sobrescribir_valor_previo": True,
        })
        aplicados += 1
    return resultado, aplicados
