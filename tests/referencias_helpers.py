"""Libros Excel sintéticos para probar la biblioteca de referencias sin depender de archivos reales."""

from __future__ import annotations

from io import BytesIO
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Border, Side


DATOS_EMPRESA = {
    "razon_social": "Acme Servicios SAS",
    "nit": "900123456",
    "direccion": "Calle 10 # 20-30",
    "telefono": "6015550101",
    "correo": "contacto@acme.test",
}


def _libro(filas: List[List[str]], titulo: str = "DATOS DEL PROVEEDOR") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Formulario"
    ws["A1"] = titulo
    lado = Side(style="thin")
    for posicion, fila in enumerate(filas):
        numero = 3 + 2 * posicion  # fila en blanco entre rótulos, como en un formulario real
        for indice in (1, 2):
            celda = ws.cell(row=numero, column=indice)
            celda.border = Border(left=lado, right=lado, top=lado, bottom=lado)
            if indice <= len(fila) and fila[indice - 1]:
                celda.value = fila[indice - 1]
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _diligenciado() -> bytes:
    return _libro([
        ["Razón social", "Acme Servicios SAS"],
        ["NIT", "900123456"],
        ["Dirección", "Calle 10 # 20-30"],
        ["Teléfono", "6015550101"],
        ["Correo electrónico", "contacto@acme.test"],
    ])


def _en_blanco() -> bytes:
    return _libro([["Razón social"], ["NIT"], ["Dirección"], ["Observaciones adicionales"]])
