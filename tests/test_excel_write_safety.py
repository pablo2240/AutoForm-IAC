"""Regresiones de integridad para el ejecutor estricto de Excel."""

from io import BytesIO
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from core.excel_inspector import inspeccionar_libro_excel
from core.excel_parser import resolver_destino_definitivo
from core.excel_verifier import verificar_integridad_excel
from core.excel_writer import rellenar_formulario_excel
from pipeline.context import PipelineContext
from pipeline.stages.stage_1_parser import ejecutar_stage_1_parser
from pipeline.stages.stage_3_llm_mapper import _expandir_listas_societarias


def _libro_base() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Formulario"
    ws["A1"] = "Razón social"
    ws["C1"] = "=LEN(B1)"
    ws.merge_cells("E1:F1")
    ws["E1"] = "Campo combinado"
    validacion = DataValidation(type="list", formula1='"SI,NO"')
    ws.add_data_validation(validacion)
    validacion.add(ws["B2"])
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def test_writer_only_modifies_the_resolved_empty_destination() -> None:
    original = _libro_base()
    plan = [{
        "hoja": "Formulario", "fila": 1, "columna": 1,
        "fila_destino": 1, "columna_destino": 2,
        "destino_resuelto": True, "campo": "razon_social", "estado": "APROBADO",
    }]
    generado, reporte = rellenar_formulario_excel(original, plan, {"razon_social": "IAC Latam"})
    assert reporte[0]["estado"] == "OK"
    wb = load_workbook(BytesIO(generado), data_only=False)
    assert wb["Formulario"]["B1"].value == "IAC Latam"
    assert wb["Formulario"]["C1"].value == "=LEN(B1)"
    assert "E1:F1" in {str(r) for r in wb["Formulario"].merged_cells.ranges}

    verificacion = verificar_integridad_excel(
        original, generado, plan, inspeccionar_libro_excel(original), reporte
    )
    assert verificacion.es_valido, verificacion.errores_bloqueantes


def test_writer_blocks_formula_and_existing_value_without_redirecting() -> None:
    original = _libro_base()
    plan = [{
        "hoja": "Formulario", "fila": 1, "columna": 1,
        "fila_destino": 1, "columna_destino": 3,
        "destino_resuelto": True, "campo": "razon_social", "estado": "APROBADO",
    }]
    generado, reporte = rellenar_formulario_excel(original, plan, {"razon_social": "IAC Latam"})
    assert reporte[0]["estado"] == "BLOCKED"
    wb = load_workbook(BytesIO(generado), data_only=False)
    assert wb["Formulario"]["C1"].value == "=LEN(B1)"
    assert wb["Formulario"]["B1"].value is None


def test_resolved_destination_uses_merge_anchor() -> None:
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A1:B1")
    destino = resolver_destino_definitivo(ws, 1, 1, "derecha")
    assert destino == {"fila_destino": 1, "columna_destino": 3}


def test_societario_uses_only_existing_empty_table_rows_and_reports_excess() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Socios"
    ws.append(["Nombre", "Identificación"])
    ws.append([None, None])
    ws.append([None, None])
    tabla = Table(displayName="SociosTabla", ref="A1:B3")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabla)
    salida = BytesIO()
    wb.save(salida)
    perfil = {"societario": {"accionistas": [
        {"nombre": "Uno", "identificacion": "1"},
        {"nombre": "Dos", "identificacion": "2"},
        {"nombre": "Tres", "identificacion": "3"},
    ]}}
    ctx = PipelineContext(archivo_bytes=salida.getvalue(), datos_empresa=perfil)
    ctx.inspeccion_excel = inspeccionar_libro_excel(ctx.archivo_bytes)
    plan = [
        {"hoja": "Socios", "fila": 1, "columna": 1, "fila_destino": 2, "columna_destino": 1, "campo": "accionista_nombre"},
        {"hoja": "Socios", "fila": 1, "columna": 2, "fila_destino": 2, "columna_destino": 2, "campo": "accionista_identificacion"},
    ]
    expandido = _expandir_listas_societarias(plan, ctx)
    assert {(p["fila_destino"], p["valor_a_escribir"]) for p in expandido if p["campo"] == "accionista_nombre"} == {(2, "Uno"), (3, "Dos")}
    assert ctx.metadatos["excedentes_no_asignados"] == [{"nombre": "Tres", "identificacion": "3"}]


def _agregar_parte_zip(libro: bytes, nombre: str, contenido: bytes) -> bytes:
    """Añade una parte opaca al paquete OOXML para probar preservación binaria."""
    salida = BytesIO()
    with zipfile.ZipFile(BytesIO(libro), "r") as origen, zipfile.ZipFile(salida, "w") as destino:
        for info in origen.infolist():
            destino.writestr(info, origen.read(info.filename))
        destino.writestr(nombre, contenido)
    return salida.getvalue()


def test_writer_preserves_vba_part_for_xlsm() -> None:
    original = _agregar_parte_zip(_libro_base(), "xl/vbaProject.bin", b"opaque-vba-project")
    plan = [{
        "hoja": "Formulario", "fila_destino": 1, "columna_destino": 2,
        "destino_resuelto": True, "campo": "razon_social", "estado": "APROBADO",
    }]
    generado, reporte = rellenar_formulario_excel(
        original, plan, {"razon_social": "IAC Latam"}, keep_vba=True
    )
    assert reporte[0]["estado"] == "OK"
    with zipfile.ZipFile(BytesIO(generado)) as archivo:
        assert archivo.read("xl/vbaProject.bin") == b"opaque-vba-project"


def test_stage_1_rejects_legacy_xls_safely() -> None:
    ctx = PipelineContext(archivo_bytes=b"legacy-xls", nombre_archivo="formulario.xls")
    with pytest.raises(ValueError, match=r"\.xls"):
        ejecutar_stage_1_parser(ctx)


def test_stage_1_allows_vml_controls_but_warns_that_they_are_not_kept() -> None:
    original = _agregar_parte_zip(_libro_base(), "xl/ctrlProps/ctrlProp1.xml", b"<ctrlProp/>")
    ctx = PipelineContext(archivo_bytes=original, nombre_archivo="formulario.xlsx")

    ctx = ejecutar_stage_1_parser(ctx)

    assert ctx.elementos_raw
    assert "controles de formulario" in ctx.metadatos["advertencia_controles_vml"]


def test_writer_processes_workbooks_with_vml_controls_without_raising() -> None:
    original = _agregar_parte_zip(_libro_base(), "xl/ctrlProps/ctrlProp1.xml", b"<ctrlProp/>")

    salida, _ = rellenar_formulario_excel(original, [], {})

    assert load_workbook(BytesIO(salida)).active.title == "Formulario"


def test_writer_keeps_leading_zeros_in_identifiers_as_text() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Formulario"
    ws["A1"] = "Número de Cuenta"
    salida = BytesIO()
    wb.save(salida)
    plan = [{
        "hoja": "Formulario", "fila": 1, "columna": 1, "fila_destino": 1, "columna_destino": 2,
        "ubicacion": "derecha", "campo": "numero_cuenta", "valor_a_escribir": "00300833888",
    }]

    resultado, _ = rellenar_formulario_excel(salida.getvalue(), plan, {})

    celda = load_workbook(BytesIO(resultado)).active["B1"]
    assert celda.value == "00300833888" and celda.data_type == "s" and celda.number_format == "@"


def _libro_con_rango_combinado() -> bytes:
    """Fila 37 con rótulos sueltos y fila 38 con una celda de captura combinada C38:H38 (como en DE-GCS-038)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "FORMATO"
    ws["F37"] = "Rut"
    ws.merge_cells("C38:H38")
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def test_un_destino_dentro_de_un_rango_combinado_se_escribe_en_el_ancla_y_se_verifica_alli() -> None:
    original = _libro_con_rango_combinado()
    plan = [{
        "hoja": "FORMATO", "fila": 37, "columna": 6, "fila_destino": 38, "columna_destino": 6,  # F38, dentro de C38:H38
        "ubicacion": "abajo", "campo": "nit", "valor": "8110047212",
    }]

    resultado, reporte = rellenar_formulario_excel(original, plan, {})
    verificacion = verificar_integridad_excel(
        original, resultado, plan, inspeccionar_libro_excel(original), reporte
    )

    hoja = load_workbook(BytesIO(resultado)).active
    assert hoja["C38"].value == "8110047212"  # el valor vive en el ancla del rango
    ok = [r for r in reporte if r["estado"] == "OK"]
    assert (ok[0]["fila_destino"], ok[0]["columna_destino"]) == (38, 3)  # se informa la celda realmente escrita
    assert verificacion.es_valido and not verificacion.errores_bloqueantes


def test_el_verificador_resuelve_el_ancla_aunque_el_reporte_traiga_otra_celda_del_rango() -> None:
    original = _libro_con_rango_combinado()
    escrito = load_workbook(BytesIO(original))
    escrito.active["C38"] = "8110047212"
    salida = BytesIO()
    escrito.save(salida)
    plan = [{"hoja": "FORMATO", "fila_destino": 38, "columna_destino": 6, "valor": "8110047212", "campo": "nit"}]

    verificacion = verificar_integridad_excel(original, salida.getvalue(), plan, inspeccionar_libro_excel(original))

    assert verificacion.es_valido and verificacion.exitosos == 1
