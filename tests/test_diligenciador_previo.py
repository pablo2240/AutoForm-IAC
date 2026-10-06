"""Reemplazo de datos de un diligenciador anterior ya escritos en el formulario (cambio de perfil)."""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from typing import Any, Dict, List

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from core.diligenciador_previo import aplicar_al_plan, detectar_reemplazos
from core.excel_writer import rellenar_formulario_excel
from pipeline.stages.stage_3d_diligenciador_previo import ejecutar_stage_3d_diligenciador_previo
from ui.page_verify import aplicar_cambios_verificacion, preparar_tabla_verificacion

PABLO = {
    "nombre": "Pablo Reyes", "cargo": "Aprendiz", "cedula": "1000000001", "telefono": "3113284008",
    "correo": "pablo.reyes@iaclatam.com", "direccion": "Calle 99 # 1-2", "ciudad": "Cali",
}
JOSE = {
    "nombre": "José Pérez", "cargo": "Comercial", "cedula": "2000000002", "telefono": "3100000002",
    "correo": "jose.perez@iaclatam.com", "direccion": "Carrera 7 # 8-9", "ciudad": "Bogotá",
}
EMPRESA = {
    "razon_social": "Acme Servicios SAS", "nit": "900123456", "direccion": "Calle 10 # 20-30",
    "ciudad": "Medellín", "telefono": "6015550101", "correo": "contacto@acme.test",
}


def _perfil_activo(operador: Dict[str, str]) -> Dict[str, Any]:
    """Datos fusionados como los entrega fusionar_operador_en_datos_empresa."""
    return {
        **EMPRESA,
        "responsable_nombre": operador["nombre"], "responsable_cargo": operador["cargo"],
        "responsable_cedula": operador["cedula"], "responsable_telefono": operador["telefono"],
        "responsable_correo": operador["correo"], "responsable_direccion": operador["direccion"],
        "responsable_ciudad": operador["ciudad"], "operador": dict(operador),
    }


def _formulario_con(operador: Dict[str, str], extra: Dict[str, Any] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Formato"
    ws["A1"] = "CONTACTO COMERCIAL"
    bloque = [
        ("Nombre", operador["nombre"]), ("Cargo", operador["cargo"]), ("Cédula", int(operador["cedula"])),
        ("Teléfono", operador["telefono"]), ("Correo", operador["correo"] + "   "),
    ]
    for i, (rotulo, valor) in enumerate(bloque):
        ws.cell(row=3 + i, column=1, value=rotulo)
        ws.cell(row=3 + i, column=2, value=valor)
    ws["A12"] = "Razón social"
    ws["B12"] = EMPRESA["razon_social"]
    for celda, valor in (extra or {}).items():
        ws[celda] = valor
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _por_campo(reemplazos: List[Any]) -> Dict[str, Any]:
    return {r.campo: r for r in reemplazos}


def test_reconoce_los_datos_del_diligenciador_anterior_y_los_asigna_al_activo() -> None:
    reemplazos = detectar_reemplazos(_formulario_con(PABLO), _perfil_activo(JOSE), [PABLO, JOSE], _perfil_activo(JOSE))

    campos = _por_campo(reemplazos)
    assert set(campos) == {
        "responsable_nombre", "responsable_cargo", "responsable_cedula", "responsable_telefono", "responsable_correo",
    }
    assert (campos["responsable_nombre"].valor_anterior, campos["responsable_nombre"].valor_nuevo) == ("Pablo Reyes", "José Pérez")
    assert campos["responsable_correo"].valor_nuevo == "jose.perez@iaclatam.com"
    assert campos["responsable_cedula"].valor_nuevo == "2000000002"
    assert campos["responsable_cargo"].valor_nuevo == "Comercial"
    assert campos["responsable_nombre"].rotulo_cercano == "Nombre"
    assert all(r.operador_previo == "Pablo Reyes" for r in reemplazos)


def test_no_toca_datos_de_la_empresa_ni_cambia_nada_si_el_activo_ya_es_el_del_formulario() -> None:
    assert detectar_reemplazos(_formulario_con(PABLO), _perfil_activo(PABLO), [PABLO], _perfil_activo(PABLO)) == []
    celdas = {(r.fila, r.columna) for r in detectar_reemplazos(_formulario_con(PABLO), _perfil_activo(JOSE), [PABLO], _perfil_activo(JOSE))}
    assert (12, 2) not in celdas  # la razón social de la empresa no se toca


def test_un_valor_que_tambien_es_dato_de_la_empresa_no_se_considera_del_diligenciador() -> None:
    previo = {**PABLO, "correo": EMPRESA["correo"], "telefono": EMPRESA["telefono"]}

    reemplazos = detectar_reemplazos(_formulario_con(previo), _perfil_activo(JOSE), [previo], _perfil_activo(JOSE))

    assert "responsable_correo" not in _por_campo(reemplazos)
    assert "responsable_telefono" not in _por_campo(reemplazos)
    assert "responsable_nombre" in _por_campo(reemplazos)


def test_el_cargo_solo_se_reemplaza_junto_a_la_identidad_del_mismo_diligenciador() -> None:
    lejos = _formulario_con(PABLO, {"A40": "Cargo", "B40": "Aprendiz"})

    reemplazos = detectar_reemplazos(lejos, _perfil_activo(JOSE), [PABLO], _perfil_activo(JOSE))

    assert all(r.fila != 40 for r in reemplazos)


def test_las_formulas_y_los_libros_sin_diligenciadores_conocidos_se_ignoran() -> None:
    assert detectar_reemplazos(_formulario_con(PABLO), _perfil_activo(JOSE), [], _perfil_activo(JOSE)) == []
    con_formula = _formulario_con(PABLO, {"D3": "=B3"})
    assert all(r.columna != 4 for r in detectar_reemplazos(con_formula, _perfil_activo(JOSE), [PABLO], _perfil_activo(JOSE)))


def test_si_el_perfil_activo_no_tiene_un_dato_el_del_anterior_se_retira() -> None:
    sin_cedula = {**JOSE, "cedula": ""}

    reemplazos = detectar_reemplazos(_formulario_con(PABLO), _perfil_activo(sin_cedula), [PABLO], _perfil_activo(sin_cedula))

    assert _por_campo(reemplazos)["responsable_cedula"].valor_nuevo == ""


def test_aplicar_al_plan_agrega_marca_o_respeta_segun_lo_que_ya_haya_en_la_celda() -> None:
    reemplazos = detectar_reemplazos(_formulario_con(PABLO), _perfil_activo(JOSE), [PABLO], _perfil_activo(JOSE))
    plan = [
        {"hoja": "Formato", "fila": 3, "columna": 1, "ubicacion": "derecha", "campo": "responsable_nombre"},   # B3
        {"hoja": "Formato", "fila": 5, "columna": 1, "ubicacion": "derecha", "campo": "representante_legal"},  # B5
    ]

    nuevo_plan, aplicados = aplicar_al_plan(plan, reemplazos)

    assert nuevo_plan[0]["sobrescribir_valor_previo"] is True and nuevo_plan[0]["campo"] == "responsable_nombre"
    assert "sobrescribir_valor_previo" not in nuevo_plan[1]  # otro dominio: no se toca
    agregados = {(i["fila_destino"], i["columna_destino"]): i for i in nuevo_plan[2:]}
    assert (4, 2) in agregados and (6, 2) in agregados and (5, 2) not in agregados
    assert all(i["sobrescribir_valor_previo"] and i["ubicacion"] == "misma" for i in agregados.values())
    assert aplicados == len(reemplazos) - 1


def test_el_escritor_sobrescribe_solo_si_el_item_esta_marcado_y_es_del_diligenciador() -> None:
    original = _formulario_con(PABLO)
    base = {"hoja": "Formato", "fila": 3, "columna": 2, "fila_destino": 3, "columna_destino": 2, "ubicacion": "misma"}
    plan = [
        {**base, "campo": "responsable_nombre", "valor_a_escribir": "José Pérez", "sobrescribir_valor_previo": True},
        {**base, "fila": 4, "fila_destino": 4, "campo": "responsable_cargo", "valor_a_escribir": "Comercial"},  # sin marca
        {**base, "fila": 12, "fila_destino": 12, "campo": "razon_social", "valor_a_escribir": "Otra SAS", "sobrescribir_valor_previo": True},
        {**base, "fila": 5, "fila_destino": 5, "campo": "responsable_cedula", "valor_a_escribir": "", "sobrescribir_valor_previo": True},
    ]

    salida, reporte = rellenar_formulario_excel(original, plan, {})

    hoja = load_workbook(BytesIO(salida)).active
    assert hoja["B3"].value == "José Pérez"
    assert hoja["B4"].value == "Aprendiz"  # sin marca: se respeta lo que ya había
    assert hoja["B12"].value == EMPRESA["razon_social"]  # un campo de la empresa nunca se sobrescribe
    assert hoja["B5"].value is None  # sin dato en el perfil activo: se retira el del anterior
    assert [r["estado"] for r in reporte].count("BLOCKED") == 2


def test_la_marca_sobrevive_a_la_tabla_de_verificacion() -> None:
    plan = [{
        "hoja": "Formato", "fila": 3, "columna": 2, "fila_destino": 3, "columna_destino": 2, "ubicacion": "misma",
        "campo": "responsable_nombre", "rotulo": "Nombre", "sobrescribir_valor_previo": True,
        "valor_anterior": "Pablo Reyes",
    }, {
        "hoja": "Formato", "fila": 12, "columna": 1, "ubicacion": "derecha", "campo": "razon_social", "rotulo": "Razón social",
    }]
    datos = _perfil_activo(JOSE)

    tabla = preparar_tabla_verificacion(plan, datos)
    resultado = aplicar_cambios_verificacion(tabla, plan, datos)

    por_campo = {i["campo"]: i for i in resultado}
    assert por_campo["responsable_nombre"]["sobrescribir_valor_previo"] is True
    assert por_campo["responsable_nombre"]["valor_anterior"] == "Pablo Reyes"
    assert "sobrescribir_valor_previo" not in por_campo["razon_social"]


def _ctx(archivo: bytes, datos: Dict[str, Any], conocidos: List[Dict[str, str]], plan: List[Dict[str, Any]] | None = None) -> Any:
    registro: List[str] = []
    return SimpleNamespace(
        archivo_bytes=archivo, datos_empresa=datos, operadores_conocidos=conocidos, tipo_documento="excel",
        plan_mapeo=list(plan or []), metadatos={}, log=lambda mensaje, *_a, **_k: registro.append(mensaje), registro=registro,
    )


def test_la_etapa_3d_deja_el_plan_listo_y_el_formulario_final_con_los_datos_del_perfil_activo() -> None:
    ctx = _ctx(_formulario_con(PABLO), _perfil_activo(JOSE), [PABLO, JOSE])

    ejecutar_stage_3d_diligenciador_previo(ctx)
    salida, _ = rellenar_formulario_excel(ctx.archivo_bytes, ctx.plan_mapeo, ctx.datos_empresa)

    hoja = load_workbook(BytesIO(salida)).active
    assert [hoja.cell(row=r, column=2).value for r in range(3, 8)] == [
        "José Pérez", "Comercial", "2000000002", "3100000002", "jose.perez@iaclatam.com",
    ]
    assert hoja["B12"].value == EMPRESA["razon_social"]
    assert ctx.metadatos["diligenciador_previo"]["operadores"] == ["Pablo Reyes"]


def test_cambiar_de_perfil_dos_veces_sobre_el_mismo_formulario_sigue_funcionando() -> None:
    primero = _ctx(_formulario_con(PABLO), _perfil_activo(JOSE), [PABLO, JOSE])
    ejecutar_stage_3d_diligenciador_previo(primero)
    salida, _ = rellenar_formulario_excel(primero.archivo_bytes, primero.plan_mapeo, primero.datos_empresa)

    segundo = _ctx(salida, _perfil_activo(PABLO), [PABLO, JOSE])
    ejecutar_stage_3d_diligenciador_previo(segundo)
    final, _ = rellenar_formulario_excel(segundo.archivo_bytes, segundo.plan_mapeo, segundo.datos_empresa)

    assert load_workbook(BytesIO(final)).active["B3"].value == "Pablo Reyes"


@pytest.mark.parametrize("datos,conocidos", [
    ({**EMPRESA}, [PABLO]),                   # sin diligenciador activo
    (_perfil_activo(JOSE), []),               # sin diligenciadores conocidos
])
def test_la_etapa_3d_no_hace_nada_sin_diligenciador_activo_o_conocidos(datos: Dict[str, Any], conocidos: List[Dict[str, str]]) -> None:
    ctx = _ctx(_formulario_con(PABLO), datos, conocidos)

    ejecutar_stage_3d_diligenciador_previo(ctx)

    assert ctx.plan_mapeo == [] and "diligenciador_previo" not in ctx.metadatos


def test_un_error_en_la_etapa_3d_nunca_impide_procesar_el_formulario() -> None:
    ctx = _ctx(b"no es un excel", _perfil_activo(JOSE), [PABLO])

    ejecutar_stage_3d_diligenciador_previo(ctx)

    assert ctx.plan_mapeo == []
    assert any("Omitida" in m for m in ctx.registro)


def test_un_correo_suelto_de_la_plantilla_no_se_confunde_con_un_diligenciador_anterior() -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "E-mail autorizado facturación"
    ws["B1"] = PABLO["correo"]
    salida = BytesIO()
    wb.save(salida)

    assert detectar_reemplazos(salida.getvalue(), _perfil_activo(JOSE), [PABLO], _perfil_activo(JOSE)) == []


def test_un_nombre_completo_basta_para_reconocer_al_diligenciador_anterior() -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Contacto"
    ws["B1"] = PABLO["nombre"]
    salida = BytesIO()
    wb.save(salida)

    reemplazos = detectar_reemplazos(salida.getvalue(), _perfil_activo(JOSE), [PABLO], _perfil_activo(JOSE))

    assert [(r.campo, r.valor_nuevo) for r in reemplazos] == [("responsable_nombre", "José Pérez")]
