"""El barrio es un dato fijo de la empresa: llega al LLM, se reconoce por su rótulo y se escribe en el formulario."""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core import profile_manager
from core.domain_constants import CAMPOS_EMPRESA, resolver_campo_por_alias_determinista
from core.excel_writer import rellenar_formulario_excel
from core.fastembed_matcher import DESCRIPCIONES_TAXONOMIA
from core.llm_client import STRICT_SYSTEM_PROMPT

RAIZ = Path(__file__).resolve().parent.parent


def test_el_perfil_de_la_empresa_incluye_el_barrio() -> None:
    datos = json.loads((RAIZ / "config" / "datos_empresa.json").read_text(encoding="utf-8-sig"))

    assert profile_manager.aplanar_perfil(datos)["barrio"] == "Los Conquistadores"


def test_la_taxonomia_coloca_el_barrio_en_la_ubicacion_de_la_empresa_sin_duplicarlo() -> None:
    taxonomia = profile_manager.estructurar_perfil_taxonomia({"razon_social": "Acme", "barrio": "Los Conquistadores"})

    assert taxonomia["empresa"]["ubicacion"]["barrio"] == "Los Conquistadores"
    assert "barrio" not in taxonomia  # no se repite como atributo suelto


def test_la_plantilla_vacia_y_el_dominio_conocen_el_barrio() -> None:
    assert profile_manager._obtener_plantilla_vacia()["barrio"] == ""
    assert "barrio" in CAMPOS_EMPRESA
    assert "barrio" in DESCRIPCIONES_TAXONOMIA


def test_el_prompt_del_llm_describe_el_barrio_de_la_empresa() -> None:
    assert "`barrio`" in STRICT_SYSTEM_PROMPT


def test_el_rotulo_barrio_se_resuelve_de_forma_determinista() -> None:
    assert resolver_campo_por_alias_determinista("Barrio") == "barrio"
    assert resolver_campo_por_alias_determinista("BARRIO:") == "barrio"
    assert resolver_campo_por_alias_determinista("Barrio / Sector") == "barrio"


def test_el_barrio_se_escribe_en_el_formulario() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Formulario"
    ws["A1"] = "Barrio"
    salida = BytesIO()
    wb.save(salida)
    plan = [{
        "hoja": "Formulario", "fila": 1, "columna": 1, "fila_destino": 1, "columna_destino": 2,
        "ubicacion": "derecha", "campo": "barrio",
    }]

    resultado, _ = rellenar_formulario_excel(
        salida.getvalue(), plan, profile_manager.estructurar_perfil_taxonomia({"barrio": "Los Conquistadores"})
    )

    assert load_workbook(BytesIO(resultado)).active["B1"].value == "Los Conquistadores"
