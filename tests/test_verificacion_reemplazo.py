"""La vista de verificación permite pedir de forma explícita que un valor reemplace lo que la celda ya trae."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

from core.excel_writer import rellenar_formulario_excel
from ui.page_verify import aplicar_cambios_verificacion, preparar_tabla_verificacion, tiene_valor_existente

DATOS = {"total_activos": "16151175009", "razon_social": "Acme SAS"}


def _plan() -> list:
    return [
        {"hoja": "Formulario", "fila": 1, "columna": 1, "fila_destino": 1, "columna_destino": 2, "ubicacion": "derecha",
         "campo": "total_activos", "rotulo": "Total activos", "valor_anterior": "20980260950"},
        {"hoja": "Formulario", "fila": 2, "columna": 1, "fila_destino": 2, "columna_destino": 2, "ubicacion": "derecha",
         "campo": "razon_social", "rotulo": "Razón social"},
    ]


def test_detecta_si_el_destino_ya_trae_contenido() -> None:
    assert tiene_valor_existente("20980260950") and tiene_valor_existente("Antonio Prieto")
    assert not any(tiene_valor_existente(v) for v in ("-", "", None, "nan", "None", "  "))


def test_la_tabla_trae_la_columna_sin_marcar_y_el_plan_automatico_no_reemplaza() -> None:
    tabla = preparar_tabla_verificacion(_plan(), DATOS)

    assert "Reemplazar existente" in tabla.columns and not tabla["Reemplazar existente"].any()
    assert all("sobrescribir_confirmado_por_usuario" not in i for i in aplicar_cambios_verificacion(tabla, _plan(), DATOS))


def test_la_marca_del_usuario_viaja_al_plan_y_el_formulario_final_refleja_el_reemplazo() -> None:
    tabla = preparar_tabla_verificacion(_plan(), DATOS)
    indice = tabla.index[tabla["Campo Asignado"] == "total_activos"][0]
    tabla.at[indice, "Reemplazar existente"] = True

    plan = aplicar_cambios_verificacion(tabla, _plan(), DATOS)
    por_campo = {i["campo"]: i for i in plan}
    assert por_campo["total_activos"]["sobrescribir_confirmado_por_usuario"] is True
    assert "sobrescribir_confirmado_por_usuario" not in por_campo["razon_social"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Formulario"
    ws["A1"], ws["B1"] = "Total activos", 20980260950
    salida = BytesIO()
    wb.save(salida)
    resultado, _ = rellenar_formulario_excel(salida.getvalue(), plan, DATOS)

    assert load_workbook(BytesIO(resultado)).active["B1"].value == "16151175009"


def test_una_sugerencia_sobre_una_celda_con_marca_de_casilla_se_omite() -> None:
    from types import SimpleNamespace

    elementos = [{
        "hoja": "Formulario", "fila": 8, "columna": 12, "valor": "PROVEEDOR", "tipo_clasificacion": "CAMPO_ENTRADA",
        "tipoEspacioEscritura": "derecha", "seccion_padre": "INFORMACIÓN GENERAL",
    }]
    inspeccion = SimpleNamespace(
        obtener_valor_anterior=lambda hoja, fila, col: "X", es_celda_formula=lambda *a: False,
        es_celda_protegida=lambda *a: False,
    )

    tabla = preparar_tabla_verificacion([], {"razon_social": "Acme SAS"}, elementos, inspeccion_excel=inspeccion)

    fila = tabla[tabla["Rótulo / Encabezado"] == "PROVEEDOR"].iloc[0]
    assert fila["Campo Asignado"].startswith("--") and "marca de casilla" in fila["Motivo / Observación"]
