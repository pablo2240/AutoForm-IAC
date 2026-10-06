"""Componente de Verificación Visual con Dropdowns para AutoForm AI (Fase 4 HSP).

Muestra TODOS los rótulos detectados en el formulario estructurados con la jerarquía
de secciones, semáforos de confianza (🟢/🟡/⚪), y motivos de validación determinística.
Permite búsqueda, filtrado por sección y estado, edición interactiva y guardado de plantillas.
"""

from __future__ import annotations

import re
import unicodedata
from typing import Any, Dict, List, Optional, Set, Tuple
import pandas as pd
import streamlit as st

try:
    from rapidfuzz import fuzz
except ImportError:
    fuzz = None

from openpyxl.utils import get_column_letter
from core.plausibilidad import evaluar_plausibilidad
from pipeline.context import PipelineContext
from template_store.store import guardar_plantilla, calcular_hash_formulario


OPCION_OMITIR = "-- Omitir / Dejar vacío --"

CAMPOS_VIRTUALES = [
    "ciudad_departamento",
    "representante_nombres",
    "representante_apellidos",
]

# Sinónimos robustos para sugerir coincidencias parciales
_SINONIMOS_RAPIDOS = {
    "razon_social": [
        "razon social", "empresa", "proveedor", "denominacion social", "sociedad", "solicitante",
        "nombre comercial", "nombre o razon social", "razon social / nombre comercial",
        "razon social o nombre comercial", "razon social nombre comercial",
    ],
    "nit": [
        "nit", "rut", "identificacion tributaria", "registro fiscal", "numero de identificacion tributaria",
        "nit tax id", "tax id", "tax id nit", "tax identification number", "nit/tax id", "tax id/nit",
        "cc ce pas nit", "cc ce nit", "cc nit", "nit cc", "cc nit rut",
    ],
    "tipo_documento": [
        "tipo de documento", "tipo documento", "tipo id", "tipo de id", "tipo identificacion",
        "tipo de identificacion", "clase de documento", "tipo doc",
        "tipo de identificacion (cc-pasaporte-ce)", "tipo de identificacion cc pasaporte ce",
        "tipo de identificacion cc ce pasaporte", "tipo de documento cc ce pasaporte",
    ],
    "cedula": [
        "cedula", "cedula de ciudadania", "documento de identidad", "documento de identificacion",
        "no de documento", "numero de documento", "identificacion del representante", "identificacion",
        "numero de identificacion", "nro de identificacion", "no de identificacion", "no identificacion",
        "identificacion no", "numero identificacion", "nro identificacion",
    ],
    "representante_legal": [
        "representante legal", "apoderado", "director general", "nombre del representante",
        "razon social o nombres y apellidos", "razon social o nombres y apellidos del representante",
        "nombres y apellidos", "nombre y apellidos",
    ],
    "representante_nombres": ["nombres del representante", "primer nombre", "segundo nombre"],
    "representante_apellidos": ["apellidos del representante", "primer apellido", "segundo apellido"],
    "direccion": ["direccion", "domicilio principal", "sede principal", "direccion fiscal", "direccion de notificacion"],
    "telefono": ["telefono", "telefono fijo", "pbx institucional", "telefono corporativo"],
    "celular": [
        "celular", "telefono movil", "celular del representante", "numero de celular",
        "telefono celular", "telefono / celular", "telefono/celular", "tel/cel", "tel celular", "movil",
    ],
    "correo": ["correo", "email", "e-mail", "correo electronico", "correo institucional"],
    "pagina_web": ["pagina web", "sitio web", "portal web", "url institucional"],
    "banco": ["banco", "entidad bancaria", "institucion financiera", "nombre de la entidad financiera", "entidad financiera"],
    "numero_cuenta": ["numero de cuenta", "no de cuenta", "nro cuenta", "cuenta bancaria"],
    "tipo_cuenta": ["tipo de cuenta", "tipo cuenta", "modalidad de cuenta"],
    "sucursal": ["sucursal", "agencia bancaria", "oficina bancaria", "sucursal bancaria"],
    "ciudad": ["ciudad", "municipio", "ciudad fiscal", "ciudad de domicilio"],
    "departamento": ["departamento", "provincia"],
    "ciudad_departamento": [
        "ciudad departamento", "ciudad / departamento", "ciudad/departamento",
        "ciudad depto", "ciudad / depto", "ciudad y departamento",
        "municipio departamento", "municipio / departamento", "municipio/departamento",
    ],
    "pais": ["pais", "nacionalidad"],
    "lugar_expedicion": ["lugar de expedicion", "ciudad de expedicion", "expedida en", "municipio de expedicion", "lugar expedicion", "expedida"],
    "expedicion": ["lugar de expedicion", "ciudad de expedicion", "expedida en", "municipio de expedicion", "lugar expedicion"],
}


def _normalizar(txt: str) -> str:
    if not txt:
        return ""
    nfd = "".join(c for c in unicodedata.normalize("NFD", str(txt).lower()) if unicodedata.category(c) != "Mn")
    return re.sub(r"[\s_\.\:\-\;\,\(\)\[\]\/\\]+", " ", nfd).strip()


def _sugerir_campo_para_rotulo(rotulo: str, campos_disponibles: List[str], seccion_padre: str = "") -> Optional[str]:
    """Intenta sugerir una coincidencia parcial para un rótulo no asignado usando tokenización de palabra completa."""
    r_norm = _normalizar(rotulo)
    if not r_norm or len(r_norm) < 3:
        return None

    # Omitir palabras reservadas de opciones o preguntas
    if r_norm in ("si", "no", "s", "n", "m", "f", "ahorros", "corriente", "otro", "otra", "na", "n a"):
        return None

    # Si el rótulo solicita una FECHA (día/mes/año), NUNCA sugerir 'lugar_expedicion' ni 'expedicion'
    if "fecha" in r_norm:
        return None

    # 1. Teléfono Celular o Teléfono/Celular -> Prioridad a CELULAR
    if re.search(r"\btel[eé]fono[\s/]*celular\b|\btel[\s/]*cel\b|\bcelular\b|\bmovil\b|\bm[oó]vil\b", r_norm):
        return "celular"

    # 2. Ciudad / Departamento combinado -> ciudad_departamento ("Medellin/Antioquia")
    if re.search(r"\bciudad[\s/]+departamento\b|\bciudad[\s/]+depto\b|\bmunicipio[\s/]+departamento\b", r_norm):
        return "ciudad_departamento"

    # 3. Razón Social o Nombres y Apellidos -> representante_legal
    if "razon social" in r_norm and ("nombres" in r_norm or "apellidos" in r_norm):
        return "representante_legal"

    # 4. Nombre Comercial -> razon_social
    if "nombre comercial" in r_norm:
        return "razon_social"

    # 5. Tipo de Identificación (CC-Pasaporte-CE) -> tipo_documento
    if "tipo" in r_norm and any(k in r_norm for k in ("documento", "identificacion", "id")):
        return "tipo_documento"

    # 6. NIT / TAX ID o rótulos compuestos con NIT (CC/CE/PAS/NIT) -> 'nit'
    if re.search(r"\bnit[\s/]*tax\s*id\b|\btax\s*id\b|\bcc[\s/]*ce[\s/]*pas[\s/]*nit\b|\bcc[\s/]*ce[\s/]*nit\b|\bcc[\s/]*nit\b|\bnit[\s/]*cc\b", r_norm):
        return "nit"

    # 7. Identificación / Número de Identificación / Nro de Identificación -> 'cedula'
    if re.search(r"\bnumero\s+de\s+identificacion\b|\bnro\s+de\s+identificacion\b|\bno\s+de\s+identificacion\b|\bno\s+identificacion\b|\bidentificacion\s+no\b|\bidentificacion\b|\bdocumento\s+de\s+identidad\b", r_norm):
        if "nit" in r_norm or "tributaria" in r_norm:
            return "nit"
        return "cedula"

    # 8. En la sección del Representante Legal, teléfono / celular corresponde al móvil ('celular')
    if seccion_padre and any(k in seccion_padre.lower() for k in ("representante", "apoderado", "persona natural")):
        if re.search(r"\btel[eé]fono\b|\bcelular\b|\bmovil\b|\bcel\b|\bno\.?\s*celular\b", r_norm):
            return "celular"

    # 9. Búsqueda por sinónimos de coincidencia exacta o frase completa
    for campo, sinonimos in _SINONIMOS_RAPIDOS.items():
        if campo in campos_disponibles or campo in CAMPOS_VIRTUALES:
            for s in sinonimos:
                if s == r_norm:
                    return campo
                if len(s) >= 4 and re.search(r"\b" + re.escape(s) + r"\b", r_norm):
                    return campo

    # 10. Búsqueda difusa estricta con rapidfuzz (token_sort_ratio >= 88)
    if fuzz is not None and len(r_norm) >= 4:
        mejor_campo = None
        mejor_score = 0.0
        for campo in campos_disponibles:
            c_norm = _normalizar(campo)
            score = float(fuzz.token_sort_ratio(r_norm, c_norm))
            if score >= 88.0 and score > mejor_score:
                mejor_score = score
                mejor_campo = campo
        if mejor_campo:
            return mejor_campo

    return None


def _resolver_valor_campo(datos_empresa: Dict[str, Any], campo: str) -> str:
    """Resuelve el valor del campo empresarial, soportando campos virtuales y anidados."""
    if not campo or campo == OPCION_OMITIR:
        return ""

    from core.profile_manager import aplanar_perfil
    plano = aplanar_perfil(datos_empresa)

    if campo in ("ciudad_departamento", "ciudad/departamento", "ciudad_depto"):
        c = str(plano.get("ciudad", "")).strip()
        d = str(plano.get("departamento", "")).strip()
        if c and d:
            return f"{c}/{d}"
        return c or d

    if campo == "representante_nombres":
        if "representante_nombres" in plano and plano["representante_nombres"]:
            return str(plano["representante_nombres"])
        rep_full = str(plano.get("representante_legal", "")).strip()
        if rep_full:
            partes = rep_full.split()
            return " ".join(partes[:-2]) if len(partes) > 2 else (partes[0] if partes else rep_full)

    if campo == "representante_apellidos":
        if "representante_apellidos" in plano and plano["representante_apellidos"]:
            return str(plano["representante_apellidos"])
        rep_full = str(plano.get("representante_legal", "")).strip()
        if rep_full:
            partes = rep_full.split()
            return " ".join(partes[-2:]) if len(partes) >= 2 else ""

    if campo in plano:
        val = plano[campo]
        if not isinstance(val, dict):
            return "" if val is None else str(val)

    return ""


def obtener_opciones_campos_empresa(datos_empresa: Dict[str, Any]) -> List[str]:
    """Construye la lista ordenada de opciones seleccionables en el dropdown."""
    from core.profile_manager import aplanar_perfil
    plano = aplanar_perfil(datos_empresa)
    
    claves_limpias = {k for k in plano.keys() if "." not in k and not isinstance(plano[k], dict)}
    claves = claves_limpias | set(CAMPOS_VIRTUALES)
    lista_ordenada = sorted(list(claves))
    return [OPCION_OMITIR] + lista_ordenada


def _calcular_celda_destino(fila: int, col: int, ubicacion: str) -> Tuple[int, int, str]:
    """Calcula la coordenada (fila, columna, A1) destino a partir de la celda rótulo."""
    u = str(ubicacion or "derecha").lower()
    if u == "abajo":
        f_dest, c_dest = fila + 1, col
    elif u == "misma":
        f_dest, c_dest = fila, col
    else:
        f_dest, c_dest = fila, col + 1
    try:
        celda_str = f"{get_column_letter(c_dest)}{f_dest}"
    except Exception:
        celda_str = f"R{f_dest}C{c_dest}"
    return f_dest, c_dest, celda_str


def _formatear_badge_estado(
    estado_raw: str,
    confianza_raw: str,
    campo_select: str,
    motivo: str = "",
    confianza_score: Optional[float] = None,
) -> str:
    """Convierte el estado y confianza técnicos en un badge visual intuitivo para el usuario con score (0-1)."""
    if not campo_select or campo_select == OPCION_OMITIR:
        if estado_raw == "DESCARTADO":
            return "0.00 ⚪ Omitido"
        if estado_raw == "EXTRA":
            return "0.00 ⚫ No es campo" if "Descartado" in motivo else "0.00 ⚪ Sin asignar"
        return "0.00 ⚪ Sin asignar"

    estado_upper = str(estado_raw or "").upper()
    confianza_upper = str(confianza_raw or "").upper()

    score = 1.0
    if confianza_score is not None:
        try:
            score = max(0.0, min(1.0, float(confianza_score)))
        except (ValueError, TypeError):
            score = 1.0
    elif confianza_upper == "EXACTA":
        score = 1.0
    elif confianza_upper == "ALTA":
        score = 0.90
    elif confianza_upper == "PARCIAL":
        score = 0.65
    elif estado_upper == "REVISION":
        score = 0.50
    else:
        score = 0.80

    if estado_upper == "APROBADO":
        if confianza_upper == "EXACTA" or score >= 0.98:
            return f"{score:.2f} 🟢 Exacta"
        if "autocorr" in motivo.lower() or "corregido" in motivo.lower():
            return f"{score:.2f} 🟢 Alta (Autocorr.)"
        if confianza_upper == "ALTA" or score >= 0.75:
            return f"{score:.2f} 🟢 Alta"
        if confianza_upper == "PARCIAL" or score >= 0.50:
            return f"{score:.2f} 🟡 Parcial"
        return f"{score:.2f} 🟢 Aprobado"

    if estado_upper == "REVISION":
        return f"{score:.2f} 🟡 Requiere Revisión"

    if estado_upper == "DESCARTADO":
        return "0.00 ⚪ Omitido"

    return f"{score:.2f} ✅ Sugerido por IA"


def preparar_tabla_verificacion(
    plan_mapeo: List[Dict[str, Any]],
    datos_empresa: Dict[str, Any],
    elementos_raw: Optional[List[Dict[str, Any]]] = None,
    inspeccion_excel: Optional[Any] = None,
) -> pd.DataFrame:
    """Transforma el plan de mapeo enriquecido en un DataFrame interactivo con soporte de jerarquía y validación."""
    from pipeline.stages.stage_2_classifier import clasificar_rotulo_individual, ClasificacionElemento

    from core.profile_manager import aplanar_perfil
    plano = aplanar_perfil(datos_empresa)
    claves_disponibles = set(plano.keys()) | set(CAMPOS_VIRTUALES)
    
    filas: List[Dict[str, Any]] = []
    # Índice de celdas ya presentes en plan_mapeo
    coordenadas_mapeadas: Set[Tuple[str, int, int]] = set()

    # 1. Agregar primero los campos mapeados (enriquecidos por Fase 2 y 3)
    for idx, item in enumerate(plan_mapeo):
        hoja = str(item.get("hoja", "") or "Hoja1")
        fila = int(item.get("fila", 0) or 0)
        col = int(item.get("columna", 0) or 0)
        rotulo = str(item.get("rotulo_original") or item.get("rotulo") or item.get("encabezado") or item.get("valor") or "").strip()
        campo = str(item.get("campo", "")).strip()
        seccion = str(item.get("seccion") or item.get("seccion_padre") or "INFORMACIÓN GENERAL").strip()
        
        estado_raw = str(item.get("estado", "APROBADO")).strip()
        confianza_raw = str(item.get("nivel_confianza", "ALTA")).strip()
        conf_score = float(item.get("confianza_score") if item.get("confianza_score") is not None else (1.0 if confianza_raw == "EXACTA" else (0.85 if confianza_raw == "ALTA" else 0.5)))
        motivo = str(item.get("motivo", "")).strip()
        tipo_elem = str(item.get("tipo_elemento", "FIELD")).strip()

        ubicacion = str(item.get("ubicacion", "derecha")).lower()
        if ubicacion not in ("derecha", "abajo", "misma"):
            ubicacion = "derecha"

        f_dest = int(item.get("fila_destino") or 0)
        c_dest = int(item.get("columna_destino") or 0)
        if not f_dest or not c_dest:
            f_dest, c_dest, celda_str = _calcular_celda_destino(fila, col, ubicacion)
        else:
            try:
                celda_str = f"{get_column_letter(c_dest)}{f_dest}"
            except Exception:
                celda_str = f"R{f_dest}C{c_dest}"
        celda = str(item.get("celda") or celda_str)

        campo_select = campo if (campo and (campo in claves_disponibles or campo in datos_empresa or "." in campo)) else OPCION_OMITIR
        valor_real = _resolver_valor_campo(datos_empresa, campo_select)

        # Valor anterior en el Excel original
        v_ant = item.get("valor_anterior")
        if v_ant is None and inspeccion_excel:
            v_ant = inspeccion_excel.obtener_valor_anterior(hoja, f_dest, c_dest)
        valor_anterior_str = "-" if v_ant is None or str(v_ant).strip() == "" else str(v_ant)

        # Advertencias y detección de bloqueo (fórmulas preexistentes)
        advertencias_lista = list(item.get("advertencias") or [])
        es_bloqueante = bool(item.get("bloqueante", False))

        if inspeccion_excel:
            if inspeccion_excel.es_celda_formula(hoja, f_dest, c_dest):
                f_val = inspeccion_excel.celdas_con_formula.get((hoja, f_dest, c_dest), "")
                # Re-enrutar si la dirección alternativa está libre
                dir_alt = "abajo" if ubicacion == "derecha" else ("derecha" if ubicacion == "abajo" else None)
                re_enrutado = False
                if dir_alt:
                    f_alt, c_alt = (fila + 1, col) if dir_alt == "abajo" else (fila, col + 1)
                    if not inspeccion_excel.es_celda_formula(hoja, f_alt, c_alt) and not inspeccion_excel.es_celda_protegida(hoja, f_alt, c_alt):
                        f_dest, c_dest = f_alt, c_alt
                        ubicacion = dir_alt
                        celda = f"{get_column_letter(c_dest)}{f_dest}"
                        re_enrutado = True
                        advertencias_lista.append(f"ℹ️ Re-enrutado a {dir_alt} ({celda}) para proteger fórmula original.")
                if not re_enrutado:
                    campo_select = OPCION_OMITIR
                    valor_real = ""
                    badge = "0.00 🛡️ Preservada"
                    advertencias_lista.append(f"ℹ️ Celda {celda} calculada por fórmula ('{f_val}'); preservada intacta sin sobreescritura.")
                    es_bloqueante = False
            elif inspeccion_excel.es_celda_protegida(hoja, f_dest, c_dest):
                campo_select = OPCION_OMITIR
                valor_real = ""
                badge = "0.00 🛡️ Protegida"
                advertencias_lista.append(f"ℹ️ Celda {celda} protegida en plantilla; preservada sin sobreescritura.")
                es_bloqueante = False

        adv_display = " | ".join(advertencias_lista) if advertencias_lista else "Ninguna"

        # Badge visual con score numérico
        badge = _formatear_badge_estado(estado_raw, confianza_raw, campo_select, motivo, confianza_score=conf_score)

        motivo_display = motivo
        if not motivo_display:
            if "Exacta" in badge:
                motivo_display = "Coincidencia directa con datos maestros"
            elif "Alta" in badge:
                motivo_display = "Emparejamiento contextual validado"

        filas.append({
            "N°": idx + 1,
            "Hoja": hoja,
            "Celda / Rango": celda,
            "Rótulo / Encabezado": rotulo,
            "Valor Anterior": valor_anterior_str,
            "Campo Asignado": campo_select,
            "Valor Nuevo": valor_real,
            "Nivel de Confianza": badge,
            "Advertencias": adv_display,
            "Dirección": ubicacion,
            "Sección": seccion,
            "Ubicación": f"{hoja} (F{fila}:C{col})" if hoja else f"Fila {fila}, Col {col}",
            "Motivo / Observación": motivo_display,
            "_hoja": hoja,
            "_fila": fila,
            "_columna": col,
            "_fila_destino": f_dest,
            "_columna_destino": c_dest,
            "_bloqueante": es_bloqueante,
            "_confianza_score": conf_score,
            "_requiereMerge": bool(item.get("requiereMerge", False)),
            "_celdasAMergear": int(item.get("celdasAMergear", 1) or 1),
            "_anchoLinea": int(item.get("anchoLinea", 1) or 1),
            "_orig_idx": idx,
            "_seccion": seccion,
            "_tipo_elemento": tipo_elem,
            "_estado_raw": estado_raw,
            "_confianza_raw": confianza_raw,
            "_sobrescribir": bool(item.get("sobrescribir_valor_previo", False)),
        })
        coordenadas_mapeadas.add((hoja, fila, col))

    # 2. Agregar todos los demás elementos detectados en elementos_raw con su rol funcional
    if elementos_raw:
        n_extra = len(filas) + 1
        for elem in elementos_raw:
            hoja_e = str(elem.get("hoja", "") or "Hoja1")
            fila_e = int(elem.get("fila", 0) or 0)
            col_e = int(elem.get("columna", 0) or 0)
            rot_e = str(elem.get("valor") or elem.get("rotulo") or "").strip()
            sec_e = str(elem.get("seccion_padre") or elem.get("seccion") or "INFORMACIÓN GENERAL").strip()

            if (hoja_e, fila_e, col_e) in coordenadas_mapeadas or not rot_e:
                continue

            tipo_clasif = elem.get("tipo_clasificacion")
            if not tipo_clasif:
                tipo_clasif_enum = clasificar_rotulo_individual(rot_e, elem)
                tipo_clasif = tipo_clasif_enum.value

            r_norm = _normalizar(rot_e)

            if tipo_clasif in (
                ClasificacionElemento.TITULO_SECCION.value,
                ClasificacionElemento.OPCION_SELECCION.value,
                ClasificacionElemento.TEXTO_LEGAL.value,
                ClasificacionElemento.INSTRUCCION_TEXTO.value,
                ClasificacionElemento.CONTROL_DOCUMENTAL.value,
                ClasificacionElemento.USO_EXCLUSIVO.value,
                ClasificacionElemento.FIRMA_ESPACIO.value,
                ClasificacionElemento.NO_APLICA.value,
            ):
                campo_final = OPCION_OMITIR
                valor_final = ""
                badge = "0.00 ⚫ No es campo"
                motivo_desc = f"Descartado: clasificado como {tipo_clasif}"
                conf_val = 0.0

            elif r_norm in ("identificacion", "id", "documento", "identificacion no", "no identificacion") and not elem.get("seccion_padre"):
                campo_final = OPCION_OMITIR
                valor_final = ""
                badge = "0.50 🟡 Requiere Revisión"
                motivo_desc = "Rótulo ambiguo sin sección específica asignada"
                conf_val = 0.5

            else:
                sugerido = _sugerir_campo_para_rotulo(rot_e, claves_disponibles, seccion_padre=sec_e)
                veredicto_e = evaluar_plausibilidad(rot_e, None, str(tipo_clasif))
                if sugerido and not veredicto_e.plausible:
                    sugerido = None
                if sugerido:
                    campo_final = sugerido
                    valor_final = _resolver_valor_campo(datos_empresa, campo_final)
                    badge = "0.70 🟡 Sugerencia"
                    motivo_desc = "Sugerencia por similitud léxica"
                    conf_val = 0.70
                else:
                    campo_final = OPCION_OMITIR
                    valor_final = ""
                    badge = "0.00 ⚪ Sin asignar"
                    motivo_desc = "No se encontró dato correspondiente en el perfil"
                    conf_val = 0.0
                    if not veredicto_e.plausible:
                        motivo_desc = f"Descartado: {veredicto_e.motivo}"
                        badge = "0.00 ⚫ No es campo"

            ancho_l = int(elem.get("anchoLinea", 1) or 1)
            ubic = str(elem.get("tipoEspacioEscritura", "derecha")).lower()
            if ubic not in ("derecha", "abajo", "misma"):
                ubic = "derecha"

            f_dest_e, c_dest_e, celda_str_e = _calcular_celda_destino(fila_e, col_e, ubic)

            v_ant_e = None
            bloqueante_e = False
            advs_e = []
            if inspeccion_excel:
                v_ant_e = inspeccion_excel.obtener_valor_anterior(hoja_e, f_dest_e, c_dest_e)
                if inspeccion_excel.es_celda_formula(hoja_e, f_dest_e, c_dest_e):
                    bloqueante_e = False
                    advs_e.append(f"ℹ️ FÓRMULA: Celda {celda_str_e} contiene cálculo de la plantilla.")
                if inspeccion_excel.es_celda_protegida(hoja_e, f_dest_e, c_dest_e):
                    bloqueante_e = False
                    advs_e.append(f"ℹ️ PROTEGIDA: Celda {celda_str_e} protegida.")
            val_ant_e_str = "-" if v_ant_e is None or str(v_ant_e).strip() == "" else str(v_ant_e)

            adv_e_str = " | ".join(advs_e) if advs_e else (motivo_desc if "Descartado" in motivo_desc else "Ninguna")

            fila_dict = {
                "N°": n_extra,
                "Hoja": hoja_e,
                "Celda / Rango": celda_str_e,
                "Rótulo / Encabezado": rot_e,
                "Valor Anterior": val_ant_e_str,
                "Campo Asignado": campo_final,
                "Valor Nuevo": valor_final,
                "Nivel de Confianza": badge,
                "Advertencias": adv_e_str,
                "Dirección": ubic,
                "Sección": sec_e,
                "Ubicación": f"{hoja_e} (F{fila_e}:C{col_e})" if hoja_e else f"Fila {fila_e}, Col {col_e}",
                "Motivo / Observación": motivo_desc,
                "_hoja": hoja_e,
                "_fila": fila_e,
                "_columna": col_e,
                "_fila_destino": f_dest_e,
                "_columna_destino": c_dest_e,
                "_bloqueante": bloqueante_e,
                "_confianza_score": conf_val,
                "_requiereMerge": bool(ancho_l > 1),
                "_celdasAMergear": ancho_l,
                "_anchoLinea": ancho_l,
                "_orig_idx": None,
                "_seccion": sec_e,
                "_tipo_elemento": tipo_clasif,
                "_estado_raw": "EXTRA",
                "_confianza_raw": "SIN_COINCIDENCIA",
            }
            filas.append(fila_dict)
            coordenadas_mapeadas.add((hoja_e, fila_e, col_e))
            n_extra += 1

    df = pd.DataFrame(filas)
    if not df.empty:
        df["N°"] = df["N°"].fillna(0).astype(int)
        df["Hoja"] = df["Hoja"].fillna("Hoja1").astype(str)
        df["Celda / Rango"] = df["Celda / Rango"].fillna("").astype(str)
        df["Rótulo / Encabezado"] = df["Rótulo / Encabezado"].fillna("").astype(str)
        df["Valor Anterior"] = df["Valor Anterior"].fillna("-").astype(str)
        df["Campo Asignado"] = df["Campo Asignado"].fillna(OPCION_OMITIR).astype(str)
        df["Valor Nuevo"] = df["Valor Nuevo"].fillna("").astype(str)
        df["Nivel de Confianza"] = df["Nivel de Confianza"].fillna("0.00 ⚪ Sin asignar").astype(str)
        df["Advertencias"] = df["Advertencias"].fillna("Ninguna").astype(str)
        df["Dirección"] = df["Dirección"].fillna("derecha").astype(str)
        df["Sección"] = df["Sección"].fillna("INFORMACIÓN GENERAL").astype(str)
        df["Ubicación"] = df["Ubicación"].fillna("").astype(str)
        df["Motivo / Observación"] = df["Motivo / Observación"].fillna("").astype(str)
        df["_hoja"] = df["_hoja"].fillna("").astype(str)
        df["_fila"] = df["_fila"].fillna(0).astype(int)
        df["_columna"] = df["_columna"].fillna(0).astype(int)
        df["_fila_destino"] = df["_fila_destino"].fillna(0).astype(int)
        df["_columna_destino"] = df["_columna_destino"].fillna(0).astype(int)
        df["_bloqueante"] = df["_bloqueante"].fillna(False).astype(bool)
        df["_confianza_score"] = df["_confianza_score"].fillna(0.0).astype(float)
        df["_requiereMerge"] = df["_requiereMerge"].fillna(False).astype(bool)
        df["_celdasAMergear"] = df["_celdasAMergear"].fillna(1).astype(int)
        df["_anchoLinea"] = df["_anchoLinea"].fillna(1).astype(int)
        df["_orig_idx"] = df["_orig_idx"].fillna(-1).astype(int)
        df["_seccion"] = df["_seccion"].fillna("").astype(str)
        df["_tipo_elemento"] = df["_tipo_elemento"].fillna("FIELD").astype(str)
        df["_estado_raw"] = df["_estado_raw"].fillna("").astype(str)
        df["_confianza_raw"] = df["_confianza_raw"].fillna("").astype(str)
    return df


def aplicar_cambios_verificacion(
    df_editado: pd.DataFrame,
    plan_original: List[Dict[str, Any]],
    datos_empresa: Dict[str, Any],
) -> List[Dict[str, Any]]:
    """Convierte el DataFrame editado por el usuario de vuelta en un plan de mapeo estructurado."""
    plan_resultado: List[Dict[str, Any]] = []

    for _, row in df_editado.iterrows():
        campo_seleccionado = str(row.get("Campo Asignado", "")).strip()
        
        # Si el usuario seleccionó omitir, no incluir en la inyección
        if not campo_seleccionado or campo_seleccionado == OPCION_OMITIR:
            continue

        direccion = str(row.get("Dirección", "derecha")).lower()
        if direccion not in ("derecha", "abajo", "misma"):
            direccion = "derecha"

        hoja = str(row.get("_hoja", row.get("Hoja", "")))
        fila = int(row.get("_fila", 0))
        columna = int(row.get("_columna", 0))
        f_dest = int(row.get("_fila_destino", 0))
        c_dest = int(row.get("_columna_destino", 0))
        celda = str(row.get("Celda / Rango", row.get("Celda", "")))
        rotulo = str(row.get("Rótulo / Encabezado", row.get("Rótulo en el Formulario", "")))
        ancho_l = int(row.get("_anchoLinea", 1) or 1)
        req_merge = bool(row.get("_requiereMerge", False) or (ancho_l > 1 and direccion == "derecha"))
        seccion = str(row.get("Sección", row.get("_seccion", "")))
        valor_nuevo = str(row.get("Valor Nuevo", ""))
        valor_anterior = str(row.get("Valor Anterior", "-"))
        confianza_score = float(row.get("_confianza_score", 1.0) or 1.0)
        bloqueante = bool(row.get("_bloqueante", False))
        adv_val = str(row.get("Advertencias", ""))
        sobrescribir = row.get("_sobrescribir", False)
        sobrescribir = sobrescribir is True or str(sobrescribir) == "True"  # una fila sin la columna llega como NaN

        item_final = {
            "hoja": hoja,
            "fila": fila,
            "columna": columna,
            "fila_destino": f_dest,
            "columna_destino": c_dest,
            "celda": celda,
            "valor": rotulo,
            "ubicacion": direccion,
            "campo": campo_seleccionado,
            "valor_a_escribir": valor_nuevo,
            "valor_anterior": valor_anterior if valor_anterior != "-" else None,
            "confianza_score": confianza_score,
            "bloqueante": bloqueante,
            "advertencias": [adv_val] if adv_val and adv_val != "Ninguna" else [],
            "requiereMerge": req_merge,
            "celdasAMergear": int(row.get("_celdasAMergear", ancho_l) or ancho_l),
            "anchoLinea": ancho_l,
            "seccion": seccion,
        }
        if sobrescribir:
            item_final["sobrescribir_valor_previo"] = True

        plan_resultado.append(item_final)

    return plan_resultado


def render_pantalla_verificacion(
    ctx: PipelineContext,
    key_prefix: str = "verif_ui",
) -> Tuple[bool, List[Dict[str, Any]]]:
    """Renderiza la interfaz de verificación visual interactiva enriquecida con HSP en Streamlit."""
    plan_activo = ctx.obtener_plan_activo()
    elementos_todos = ctx.elementos_clasificados if ctx.elementos_clasificados else ctx.elementos_raw

    if not plan_activo and not elementos_todos:
        st.warning("⚠️ No se encontraron campos detectados para verificar en este formulario.")
        return False, []

    opciones_campos = obtener_opciones_campos_empresa(ctx.datos_empresa)

    st.markdown("### 📋 Vista de Cambios Propuestos y Verificación de Celdas")
    st.markdown(
        "Revisa la hoja, celda destino, valor anterior y nuevo valor propuesto antes de escribir en el archivo. "
        "El motor de precisión protege al 100% las fórmulas y validaciones preexistentes del Excel."
    )

    # ── Gestión de Estado Maestro en Session State (Aislado por Documento) ──
    import hashlib
    doc_hash = hashlib.md5(f"{ctx.nombre_archivo}_{len(ctx.archivo_bytes or b'')}".encode("utf-8", errors="ignore")).hexdigest()[:10]
    master_key = f"{key_prefix}_master_df_{doc_hash}"
    editor_key = f"{key_prefix}_data_editor_{doc_hash}"

    insp = getattr(ctx, "inspeccion_excel", None)

    if master_key not in st.session_state or st.session_state[master_key] is None:
        st.session_state[master_key] = preparar_tabla_verificacion(
            plan_activo,
            ctx.datos_empresa,
            elementos_todos,
            inspeccion_excel=insp,
        )

    master_df: pd.DataFrame = st.session_state[master_key]

    # Conteo de métricas cuantitativas
    total_detectados = len(master_df)
    total_asignados = sum(1 for c in master_df["Campo Asignado"] if c != OPCION_OMITIR)
    total_alta_confianza = sum(1 for _, r in master_df.iterrows() if "🟢" in str(r["Nivel de Confianza"]) and r["Campo Asignado"] != OPCION_OMITIR)
    total_revision = sum(1 for _, r in master_df.iterrows() if "🟡" in str(r["Nivel de Confianza"]) or (r["Campo Asignado"] == OPCION_OMITIR and "⚪ Sin asignar" in str(r["Nivel de Confianza"])))
    total_bloqueantes = sum(1 for _, r in master_df.iterrows() if r["Campo Asignado"] != OPCION_OMITIR and (bool(r.get("_bloqueante", False)) or "🚫" in str(r.get("Advertencias", ""))))
    total_editados = sum(1 for _, r in master_df.iterrows() if "✏️" in str(r["Nivel de Confianza"]))

    # ── Tarjetas Métricas Resumidas ──
    col1, col2, col3, col4 = st.columns(4)
    with col1:
        st.metric("📄 Formulario", ctx.nombre_archivo or "Documento", help=f"Total elementos analizados: {total_detectados}")
    with col2:
        st.metric("🟢 Alta Confianza", f"{total_alta_confianza}", help="Campos validados y aprobados automáticamente")
    with col3:
        st.metric("🟡 Requieren Revisión", f"{total_revision}", help="Campos con sugerencia parcial o pendientes de asignar")
    total_preservadas = sum(1 for _, r in master_df.iterrows() if "preservad" in str(r.get("Advertencias", "")).lower() or "protegid" in str(r.get("Advertencias", "")).lower())
    with col4:
        if total_preservadas > 0:
            st.metric("🛡️ Preservadas", f"{total_preservadas}", help="Celdas con fórmulas o protegidas preservadas automáticamente")
        else:
            st.metric("✏️ Asignados", f"{total_asignados}", delta=f"{total_editados} manuales" if total_editados > 0 else None)

    # ── Mensaje informativo de integridad ──
    if total_preservadas > 0:
        st.info(
            f"🛡️ **Protección de Fórmulas y Cálculos Activa**: Se preservaron {total_preservadas} celdas calculadas o protegidas "
            "intactas para garantizar la lógica original de la plantilla. El documento puede descargarse normalmente."
        )

    # ── Filtros y Búsqueda Avanzada ──
    col_search, col_sec_filter, col_filter = st.columns([2, 1.5, 1.5])
    
    with col_search:
        filtro_texto = st.text_input(
            "🔍 Buscar por rótulo, celda o campo...",
            key=f"{key_prefix}_search_input_{doc_hash}",
            placeholder="Ej: B5, NIT, Representante, Total, Razón Social...",
        ).strip().lower()

    # Obtener lista única de secciones
    secciones_disponibles = sorted(list({str(s) for s in master_df["Sección"].dropna().unique() if str(s).strip()}))
    
    with col_sec_filter:
        filtro_seccion = st.selectbox(
            "Filtrar por Sección:",
            ["Todas las secciones"] + secciones_disponibles,
            key=f"{key_prefix}_sec_filter_{doc_hash}",
        )

    with col_filter:
        vista_filtro = st.selectbox(
            "Filtrar por Estado:",
            [
                "Todos",
                "🟢 Solo Alta Confianza",
                "🟡 Solo Requieren Revisión",
                "🛡️ Celdas con Fórmulas / Protegidas",
                "✏️ Solo Editados",
                "⚪ Sin Asignar",
            ],
            key=f"{key_prefix}_vista_filter_{doc_hash}",
        )

    # Aplicar filtrado al DataFrame visual desde master_df
    df_filtrado = master_df.copy()
    
    if filtro_texto:
        df_filtrado = df_filtrado[
            df_filtrado["Rótulo / Encabezado"].astype(str).str.lower().str.contains(filtro_texto) |
            df_filtrado["Celda / Rango"].astype(str).str.lower().str.contains(filtro_texto) |
            df_filtrado["Hoja"].astype(str).str.lower().str.contains(filtro_texto) |
            df_filtrado["Campo Asignado"].astype(str).str.lower().str.contains(filtro_texto) |
            df_filtrado["Advertencias"].astype(str).str.lower().str.contains(filtro_texto) |
            df_filtrado["Sección"].astype(str).str.lower().str.contains(filtro_texto)
        ]

    if filtro_seccion != "Todas las secciones":
        df_filtrado = df_filtrado[df_filtrado["Sección"] == filtro_seccion]

    if vista_filtro == "🟢 Solo Alta Confianza":
        df_filtrado = df_filtrado[df_filtrado["Nivel de Confianza"].astype(str).str.contains("🟢")]
    elif vista_filtro == "🟡 Solo Requieren Revisión":
        df_filtrado = df_filtrado[
            df_filtrado["Nivel de Confianza"].astype(str).str.contains("🟡") |
            ((df_filtrado["Campo Asignado"] == OPCION_OMITIR) & (~df_filtrado["Nivel de Confianza"].astype(str).str.contains("⚫")))
        ]
    elif vista_filtro == "🛡️ Celdas con Fórmulas / Protegidas":
        df_filtrado = df_filtrado[
            df_filtrado["Advertencias"].astype(str).str.contains("FÓRMULA|formula|protegid", case=False)
        ]
    elif vista_filtro == "✏️ Solo Editados":
        df_filtrado = df_filtrado[df_filtrado["Nivel de Confianza"].astype(str).str.contains("✏️")]
    elif vista_filtro == "⚪ Sin Asignar":
        df_filtrado = df_filtrado[df_filtrado["Campo Asignado"] == OPCION_OMITIR]

    # Configuración de columnas interactivas para st.data_editor
    config_columnas = {
        "N°": st.column_config.NumberColumn("N°", width="small", disabled=True),
        "Hoja": st.column_config.TextColumn("Hoja", width="small", disabled=True),
        "Celda / Rango": st.column_config.TextColumn("Celda / Rango", width="small", disabled=True),
        "Rótulo / Encabezado": st.column_config.TextColumn("Rótulo / Encabezado", width="large", disabled=True),
        "Valor Anterior": st.column_config.TextColumn("Valor Anterior", width="medium", disabled=True),
        "Campo Asignado": st.column_config.SelectboxColumn(
            "Campo Asignado",
            help="Selecciona qué dato de tu perfil empresarial debe inyectarse aquí",
            width="large",
            options=opciones_campos,
            required=True,
        ),
        "Valor Nuevo": st.column_config.TextColumn("Valor Nuevo", width="medium", disabled=True),
        "Nivel de Confianza": st.column_config.TextColumn("Confianza (0-1)", width="medium", disabled=True),
        "Advertencias": st.column_config.TextColumn("Advertencias / Reglas", width="large", disabled=True),
        "Dirección": st.column_config.SelectboxColumn(
            "Dirección",
            help="Hacia dónde se inyectará el dato respecto al rótulo",
            width="small",
            options=["derecha", "abajo", "misma"],
            required=True,
        ),
        "Sección": st.column_config.TextColumn("Sección", width="medium", disabled=True),
        # Columnas internas ocultas
        "Ubicación": None,
        "Motivo / Observación": None,
        "_hoja": None,
        "_fila": None,
        "_columna": None,
        "_fila_destino": None,
        "_columna_destino": None,
        "_bloqueante": None,
        "_confianza_score": None,
        "_requiereMerge": None,
        "_celdasAMergear": None,
        "_anchoLinea": None,
        "_orig_idx": None,
        "_seccion": None,
        "_tipo_elemento": None,
        "_estado_raw": None,
        "_confianza_raw": None,
    }

    # Render del editor de datos de Streamlit
    df_editado = st.data_editor(
        df_filtrado,
        column_config=config_columnas,
        width="stretch",
        hide_index=True,
        num_rows="fixed",
        key=editor_key,
    )

    # ── Sincronizar ediciones del usuario hacia master_df ──
    hubo_cambios = False
    if df_editado is not None and not df_editado.empty:
        for _, edit_row in df_editado.iterrows():
            num_item = edit_row["N°"]
            match_idx = master_df[master_df["N°"] == num_item].index
            if not match_idx.empty:
                idx = match_idx[0]
                campo_nuevo = str(edit_row.get("Campo Asignado", "")).strip()
                dir_nueva = str(edit_row.get("Dirección", "derecha")).lower()
                
                campo_antiguo = str(master_df.at[idx, "Campo Asignado"]).strip()
                dir_antigua = str(master_df.at[idx, "Dirección"]).lower()

                if campo_nuevo != campo_antiguo or dir_nueva != dir_antigua:
                    hubo_cambios = True
                    master_df.at[idx, "Campo Asignado"] = campo_nuevo
                    master_df.at[idx, "Dirección"] = dir_nueva

                    hoja_act = str(master_df.at[idx, "_hoja"])
                    fila_act = int(master_df.at[idx, "_fila"])
                    col_act = int(master_df.at[idx, "_columna"])
                    f_dest, c_dest, celda_str = _calcular_celda_destino(fila_act, col_act, dir_nueva)
                    master_df.at[idx, "_fila_destino"] = f_dest
                    master_df.at[idx, "_columna_destino"] = c_dest
                    master_df.at[idx, "Celda / Rango"] = celda_str

                    if insp:
                        v_ant = insp.obtener_valor_anterior(hoja_act, f_dest, c_dest)
                        master_df.at[idx, "Valor Anterior"] = "-" if v_ant is None or str(v_ant).strip() == "" else str(v_ant)

                    if campo_nuevo and campo_nuevo != OPCION_OMITIR:
                        val_res = _resolver_valor_campo(ctx.datos_empresa, campo_nuevo)
                        master_df.at[idx, "Valor Nuevo"] = val_res
                        master_df.at[idx, "Nivel de Confianza"] = "1.00 ✏️ Editado por Usuario"
                        master_df.at[idx, "_confianza_score"] = 1.0

                        advs = []
                        bloquea = False
                        if insp:
                            if insp.es_celda_formula(hoja_act, f_dest, c_dest):
                                f_orig = insp.celdas_con_formula.get((hoja_act, f_dest, c_dest), "")
                                advs.append(f"ℹ️ FÓRMULA: Celda {celda_str} tiene fórmula ('{f_orig}'). Se preservará el cálculo.")
                                bloquea = False
                            if insp.es_celda_protegida(hoja_act, f_dest, c_dest):
                                advs.append(f"ℹ️ Celda {celda_str} protegida; se preservará sin sobreescribir.")
                                bloquea = False
                            reg_dv = insp.obtener_validacion(hoja_act, f_dest, c_dest)
                            if reg_dv and reg_dv.tipo == "list" and reg_dv.opciones_permitidas:
                                if not any(str(val_res).strip().lower() == op.strip().lower() for op in reg_dv.opciones_permitidas):
                                    advs.append(f"⚠️ Opciones: {reg_dv.opciones_permitidas}")
                        
                        from core.semantic_validator import validar_formato_dato
                        ok_fmt, msg_fmt = validar_formato_dato(campo_nuevo, val_res)
                        if not ok_fmt:
                            advs.append(f"⚠️ {msg_fmt}")

                        master_df.at[idx, "Advertencias"] = " | ".join(advs) if advs else "Ninguna"
                        master_df.at[idx, "_bloqueante"] = bloquea
                    else:
                        master_df.at[idx, "Valor Nuevo"] = ""
                        master_df.at[idx, "Nivel de Confianza"] = "0.00 ⚪ Omitido"
                        master_df.at[idx, "Advertencias"] = "Omitido por el usuario"
                        master_df.at[idx, "_bloqueante"] = False
                        master_df.at[idx, "_confianza_score"] = 0.0

        if hubo_cambios:
            st.session_state[master_key] = master_df

    plan_actualizado = aplicar_cambios_verificacion(master_df, plan_activo, ctx.datos_empresa)

    st.markdown("---")

    # ── Botonera de Acciones ──
    col_guardar, col_confirmar, col_reset = st.columns([1.5, 2, 1.2])

    with col_guardar:
        if st.button("💾 Guardar como Plantilla Permanente", key=f"{key_prefix}_btn_save", help="Guarda esta configuración para que futuros formularios iguales se llenen automáticamente."):
            plan_actualizado = aplicar_cambios_verificacion(master_df, plan_activo, ctx.datos_empresa)
            plantilla_id = ctx.plantilla_id or calcular_hash_formulario(ctx.elementos_raw or plan_activo)
            ruta_guardada = guardar_plantilla(
                plantilla_id=plantilla_id,
                nombre_formulario=ctx.nombre_archivo or f"Formulario_{plantilla_id[:8]}",
                tipo_documento=ctx.tipo_documento or "excel",
                elementos_raw=ctx.elementos_raw or plan_activo,
                plan_mapeo=plan_actualizado,
                metadatos={"guardado_desde_ui": True},
            )
            ctx.plantilla_id = plantilla_id
            ctx.es_plantilla_guardada = True
            st.success(f"✅ ¡Plantilla guardada permanentemente! ({len(plan_actualizado)} campos registrados en `{ruta_guardada.name}`).")

    with col_confirmar:
        confirmar_click = st.button(
            "⚡ Confirmar y Rellenar Formulario",
            type="primary",
            key=f"{key_prefix}_btn_confirm",
            help="Inyecta los datos confirmados en el archivo original y genera la descarga.",
        )
        if confirmar_click:
            plan_actualizado = aplicar_cambios_verificacion(master_df, plan_activo, ctx.datos_empresa)
            ctx.plan_verificado = plan_actualizado
            ctx.log(f"Plan de mapeo verificado y confirmado por el usuario ({len(plan_actualizado)} campos).")
            return True, plan_actualizado

    with col_reset:
        if st.button("🔄 Restablecer", key=f"{key_prefix}_btn_reset", help="Restaura las sugerencias iniciales de la IA"):
            if master_key in st.session_state:
                del st.session_state[master_key]
            ctx.plan_verificado = []
            st.rerun()

    return False, plan_actualizado
